"""Render report-only structured blocks to genuine DOCX or PDF; no HTML execution."""
import json
import re
import os
from pathlib import Path
import subprocess
import sys
import tempfile


def export_report(payload):
    kind = payload.get('format')
    if kind not in ('pdf', 'docx'):
        raise ValueError('PDF 또는 Word 형식을 선택해 주세요.')
    sections = payload.get('sections')
    if not isinstance(sections, list) or not sections or len(sections) > 10:
        raise ValueError('내보낼 의견이 없습니다.')
    runtime = Path(os.environ.get('CREDIT_REVIEW_EXPORT_PYTHON', str(Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/python/python.exe')))
    if not runtime.is_file():
        raise ValueError('문서 생성용 Python 실행 환경을 확인해 주세요.')
    with tempfile.TemporaryDirectory(prefix='credit-export-') as temp:
        source = Path(temp) / 'input.json'
        target = Path(temp) / ('report.' + kind)
        source.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
        subprocess.run([str(runtime), str(Path(__file__).resolve()), str(source), str(target)], check=True, timeout=90, capture_output=True)
        return target.read_bytes()


def render(payload, target):
    from xml.sax.saxutils import escape
    title = str(payload.get('title') or '심사의견')[:200]
    sections = payload['sections']
    def numeric(value):
        return isinstance(value,(int,float)) or isinstance(value,str) and bool(re.fullmatch(r'[-+]?\d[\d,]*(?:\.\d+)?%?',value.strip()))
    def display(value, group_digits=False):
        if value is None:return '—'
        if group_digits and isinstance(value, str):
            # Preserve headers, labels, leading-zero codes, precision and unit suffixes.
            match = re.fullmatch(r'([-+]?)([1-9]\d{4,})(\.\d+)?(%)?', value.strip())
            if match:
                sign, whole, fraction, suffix = match.groups()
                return sign + format(int(whole), ',') + (fraction or '') + (suffix or '')
        return format(value, ',') if isinstance(value,(int,float)) else str(value)
    def widths(block,count):
        raw=block.get('column_widths')
        if not isinstance(raw,list) or len(raw)!=count or not all(isinstance(x,(int,float)) and x>0 for x in raw):return [1/count]*count
        return [x/sum(raw) for x in raw]
    if payload['format'] == 'docx':
        from docx import Document
        from docx.shared import Cm, Pt
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()
        # Remove template decoration/grid rules so exported typography is predictable.
        for style in doc.styles:
            for border in list(style.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
            if style.type == 1:
                snap = OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0')
                style.element.get_or_add_pPr().append(snap)
        page = doc.sections[0]
        page.page_width, page.page_height = Cm(21), Cm(29.7)
        page.top_margin = page.bottom_margin = Cm(1.8)
        page.left_margin = page.right_margin = Cm(1.6)
        for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
            style = doc.styles[name]
            style.font.name = '맑은 고딕'
            style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            style.font.size = Pt(10 if name == 'Normal' else 13)
            from docx.shared import RGBColor
            style.font.color.rgb = RGBColor(0x30, 0x37, 0x34)
        normal = doc.styles['Normal'].paragraph_format
        normal.line_spacing = 1.5
        normal.space_after = Pt(8)
        doc.add_paragraph(title, 'Title')
        for index, section in enumerate(sections):
            if index:
                doc.add_page_break()
            doc.add_heading(str(section['title']), 1)
            for block in section['blocks']:
                if block['type'] == 'table':
                    if block.get('caption'):
                        doc.add_paragraph(str(block['caption']))
                    rows = block.get('rows', [])
                    if not rows:
                        continue
                    count = max(len(row) for row in rows)
                    if not count or count > 30:
                        continue
                    table = doc.add_table(rows=0, cols=count)
                    table.style = 'Table Grid'
                    table.autofit = False
                    for col, fraction in zip(table.columns,widths(block,count)):col.width=Cm(17.8*fraction)
                    for i, row in enumerate(rows):
                        cells = table.add_row().cells
                        for j, value in enumerate(row):
                            cells[j].text = display(value, i > 0 and j > 0)
                            from docx.enum.text import WD_ALIGN_PARAGRAPH
                            if i==0:
                                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEF3EB');cells[j]._tc.get_or_add_tcPr().append(shade)
                            for p in cells[j].paragraphs:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.RIGHT if numeric(value) else WD_ALIGN_PARAGRAPH.LEFT
                                p.paragraph_format.space_after = Pt(3)
                                p.paragraph_format.line_spacing = 1.2
                                for run in p.runs:
                                    run.font.size = Pt(9)
                                    run.bold = i == 0
                        if i == 0:
                            repeat = OxmlElement('w:tblHeader')
                            table.rows[0]._tr.get_or_add_trPr().append(repeat)
                    doc.add_paragraph()
                elif block['type'] == 'heading':
                    doc.add_heading(str(block['text']), 2)
                else:
                    doc.add_paragraph(str(block['text']))
        doc.save(target)
    else:
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, PageBreak
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        fonts = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts'
        pdfmetrics.registerFont(TTFont('Malgun', str(fonts / 'malgun.ttf')))
        pdfmetrics.registerFont(TTFont('MalgunBold', str(fonts / 'malgunbd.ttf')))
        body = ParagraphStyle('body', fontName='Malgun', fontSize=10, leading=16, spaceAfter=9, wordWrap='CJK')
        heading = ParagraphStyle('heading', parent=body, fontName='MalgunBold', fontSize=13, leading=19, spaceBefore=9, spaceAfter=9, keepWithNext=True)
        cell = ParagraphStyle('cell', parent=body, fontSize=8.5, leading=12, spaceAfter=0)
        caption = ParagraphStyle('caption', parent=body, keepWithNext=True)
        def p(text, style=body):
            return Paragraph(escape('' if text is None else str(text)).replace('\n', '<br/>'), style)
        story = [p(title, heading)]
        for index, section in enumerate(sections):
            if index:
                story.append(PageBreak())
            story.append(p(section['title'], heading))
            for block in section['blocks']:
                if block['type'] == 'table':
                    if block.get('caption'):
                        story.append(p(block['caption'],caption))
                    rows = block.get('rows', [])
                    if not rows:
                        continue
                    count = max(len(row) for row in rows)
                    if not count or count > 30:
                        continue
                    headercell=ParagraphStyle('headercell',parent=cell,fontName='MalgunBold',alignment=1)
                    numericcell=ParagraphStyle('numericcell',parent=cell,alignment=2)
                    data = [[p(display(value, i > 0 and j > 0), headercell if i==0 else numericcell if numeric(value) else cell) for j,value in enumerate(row)] + [p('', cell)] * (count - len(row)) for i,row in enumerate(rows)]
                    table = LongTable(data, colWidths=[(A4[0]-90)*v for v in widths(block,count)], repeatRows=1, splitInRow=1)
                    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.HexColor('#ccd5c8')),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#edf2e9')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
                    story.extend([table, Spacer(1, 10)])
                else:
                    story.append(p(block['text'], heading if block['type']=='heading' else body))
        def footer(canvas, doc):
            canvas.setFont('Malgun', 8)
            canvas.drawCentredString(A4[0]/2, 22, str(doc.page))
        SimpleDocTemplate(str(target), pagesize=A4, leftMargin=45, rightMargin=45, topMargin=45, bottomMargin=40, title=title, author='').build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == '__main__':
    render(json.loads(Path(sys.argv[1]).read_text(encoding='utf-8')), Path(sys.argv[2]))
