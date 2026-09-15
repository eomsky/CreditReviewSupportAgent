"""Persistent PDF/JSON sources and metadata-aware evidence selection."""
import hashlib
import json
import re
from pathlib import Path
import pymupdf as fitz

PRIORITIES = {'매우 중요':5, '중요':4, '보통':3, '낮음':2, '매우 낮음':1, '높음':4}

class DocumentError(ValueError):
    pass

class DocumentStore:
    def __init__(self, root):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.cache = {}

    def register(self, data, name):
        key = hashlib.sha256(data).hexdigest()
        name = str(name).replace('\\','/').split('/')[-1]
        suffix = Path(name).suffix.lower()
        if suffix not in ('.docx', '.pptx', '.hwp', '.hwpx', '.gif', '.bmp', '.tif', '.tiff', '.xlsx', '.xls', '.pdf', '.json', '.txt', '.csv', '.md', '.png', '.jpg', '.jpeg', '.webp'):
            raise DocumentError('DOCX·PPTX·HWP·HWPX·Excel·PDF·텍스트·이미지 파일을 사용할 수 있습니다. 구형 DOC·PPT는 DOCX·PPTX로 저장해 주세요.')
        if len(data) > 200*1024*1024:
            raise DocumentError('파일당 최대 200MB입니다.')
        if key in self.cache:
            return key
        sources = []
        if suffix in ('.png', '.jpg', '.jpeg', '.webp', '.gif', '.bmp', '.tif', '.tiff'):
            import io
            from PIL import Image
            try:
                with Image.open(io.BytesIO(data)) as image:
                    if image.width*image.height > 25000000:
                        raise ValueError('too large')
                    image.load()
                    image = image.convert('RGB')
                    image.thumbnail((2048, 2048))
                    output = io.BytesIO()
                    image.save(output, format='PNG')
                    data = output.getvalue()
            except Exception:
                raise DocumentError('이미지를 읽을 수 없습니다. 2,500만 화소 이하의 PNG·JPG·WEBP를 첨부해 주세요.') from None
            doc = {'id':key, 'name':name, 'format':'image', 'mime':'image/png', 'sources':[]}
            (self.root/(key+'.bin')).write_bytes(data)
            (self.root/(key+'.json')).write_text(json.dumps(doc,ensure_ascii=False),encoding='utf-8')
            self.cache[key] = doc
            return key
        if suffix in ('.docx','.pptx','.hwp','.hwpx'):
            from office_text import extract
            try:
                parts=extract(data,suffix)
            except Exception as error:
                raise DocumentError(f'{name}: 문서를 읽을 수 없습니다. '+(str(error) if isinstance(error,ValueError) else '파일 형식과 암호 설정을 확인해 주세요.')) from None
            for location,text in parts:
                for offset in range(0,len(text),1700):
                    sources.append({'id':f'{key}-c{len(sources)}','document_id':key,'document_name':name,'location':location,'text':location+'\n'+text[offset:offset+1700],'format':suffix[1:]})
        elif suffix in ('.xlsx','.xls'):
            import io
            try:
                if suffix=='.xlsx':
                    import openpyxl
                    values=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
                    formulas=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False)
                    sheets=[]
                    try:
                        for ws in formulas:
                            if ws.max_row is None or ws.max_column is None:ws.calculate_dimension(force=True)
                            if ws.max_row*ws.max_column>1000000:raise DocumentError('시트당 100만 셀 이하의 파일을 첨부해 주세요.')
                            rows=[]
                            for fr,vr in zip(ws.iter_rows(),values[ws.title].iter_rows()):
                                parts=[]
                                for f,v in zip(fr,vr):
                                    if f.value is None:continue
                                    text=str(v.value) if v.value is not None else ('[수식 결과 미저장: '+str(f.value)+']' if f.data_type=='f' else str(f.value))
                                    parts.append(f.coordinate+'='+text)
                                if parts:rows.append(' | '.join(parts))
                            sheets.append((ws.title,rows))
                    finally:values.close();formulas.close()
                else:
                    import xlrd
                    book=xlrd.open_workbook(file_contents=data)
                    sheets=[(ws.name,[' | '.join('R'+str(r+1)+'C'+str(c+1)+'='+str(ws.cell_value(r,c)) for c in range(ws.ncols) if ws.cell_value(r,c)!='') for r in range(ws.nrows)]) for ws in book.sheets()]
                for sheet,rows in sheets:
                    chunk=''
                    def append_chunk(text):
                        sources.append({'id':f'{key}-c{len(sources)}','document_id':key,'document_name':name,'sheet':sheet,'text':'시트: '+sheet+'\n'+text,'format':suffix[1:]})
                    for row in rows:
                        if not row:continue
                        if chunk and len(chunk)+len(row)>1800:append_chunk(chunk);chunk=''
                        for pos in range(0,len(row),1700):
                            part=row[pos:pos+1700]
                            if chunk and len(chunk)+len(part)>1800:append_chunk(chunk);chunk=''
                            chunk+=part+'\n'
                    if chunk:append_chunk(chunk)
            except DocumentError:raise
            except Exception:raise DocumentError(f'{name}: 엑셀을 읽을 수 없습니다. 암호를 해제하고 XLSX 또는 XLS 형식으로 저장해 주세요.') from None
        elif suffix == '.pdf':
            try:
                with fitz.open(stream=data, filetype='pdf') as pdf:
                    if pdf.needs_pass:
                        raise DocumentError(f'{name}: 암호를 해제한 PDF를 첨부해 주세요.')
                    for i, page in enumerate(pdf):
                        h = page.rect.height
                        for band in range(3):
                            rect = fitz.Rect(0,max(0,band*h/3-30),page.rect.width,min(h,(band+1)*h/3+30))
                            text = page.get_text('text', clip=rect, sort=True).strip()
                            if text:
                                sid = f'{key}-p{i+1}-r{band+1}'
                                sources.append({'id':sid,'document_id':key,'document_name':name,'page':i+1,'bbox':list(rect),'text':text,'image_url':f'/evidence/{key}/p{i+1}-r{band+1}.png'})
            except DocumentError:
                raise
            except Exception:
                raise DocumentError(f'{name}: PDF를 읽을 수 없습니다.') from None
        else:
            try:
                text = data.decode('utf-8-sig')
                if suffix=='.json':json.loads(text)
            except Exception:
                raise DocumentError(f'{name}: UTF-8 텍스트 또는 JSON 형식을 확인해 주세요.') from None
            # Preserve original characters, including spacing and line breaks.
            for offset in range(0,len(text),1800):
                sources.append({'id':f'{key}-c{offset}','document_id':key,'document_name':name,'text':text[offset:offset+1800], 'character_start':offset,'format':'json'})
        if not sources:
            raise DocumentError(f'{name}: 추출 가능한 본문이 없습니다. 스캔 PDF는 OCR 처리 후 첨부해 주세요.')
        doc = {'id':key,'name':name,'format':suffix[1:],'sources':sources}
        (self.root/(key+'.bin')).write_bytes(data)
        (self.root/(key+'.json')).write_text(json.dumps(doc,ensure_ascii=False),encoding='utf-8')
        self.cache[key] = doc
        return key

    def get(self, key):
        if not re.fullmatch('[0-9a-f]{64}',str(key)):
            raise DocumentError('자료 식별자가 올바르지 않습니다.')
        if key not in self.cache:
            path = self.root/(key+'.json')
            if not path.exists():
                raise DocumentError('업로드 자료를 찾을 수 없습니다. 자료를 다시 첨부해 주세요.')
            self.cache[key] = json.loads(path.read_text(encoding='utf-8'))
        return self.cache[key]

    def manifest(self, uploads):
        if not uploads:
            raise DocumentError('분석할 자료를 첨부해 주세요.')
        rows, seen = [], set()
        for item in uploads:
            key = item.get('id')
            if key in seen:
                raise DocumentError('같은 내용의 파일이 중복 첨부되어 있습니다. 하나만 남겨 주세요.')
            seen.add(key)
            doc = self.get(key)
            priority = item.get('priority') or '보통'
            if priority not in PRIORITIES:
                raise DocumentError('자료 우선순위를 다시 선택해 주세요.')
            required = item.get('required',False)
            if not isinstance(required,bool):
                raise DocumentError('반드시 포함 여부는 네 또는 아니요로 설정해 주세요.')
            rows.append({'id':key,'name':str(item.get('name') or doc['name']), 'description':str(item.get('description') or ''),'priority':priority,'required':required})
        return rows

    def select(self, terms, manifest, budget=16000, limit=24):
        ranked, by_doc = [], {}
        for meta in manifest:
            description = re.sub(r'\s+','',meta['description'])
            metadata_match = sum(term in description for term in terms)
            candidates = []
            for source in self.get(meta['id'])['sources']:
                if len(re.findall(r'\.{5,}',source['text'])) >= 3:
                    continue  # Table-of-contents leaders are not substantive evidence.
                body = re.sub(r'\s+','',source['text'])
                relevance = sum(min(body.count(term),4) for term in terms)
                score = (relevance*4+min(metadata_match,3))*PRIORITIES[meta['priority']]
                row = (score,source,meta,relevance)
                candidates.append(row)
                if score:
                    ranked.append(row)
            candidates.sort(key=lambda row:(-row[0],row[1]['id']))
            by_doc[meta['id']] = candidates
        chosen, ids, used = [], set(), 0
        def add(row, mandatory=False):
            nonlocal used
            _, source, meta, relevance = row
            if source['id'] in ids:
                return True
            if used+len(source['text'])>budget or len(chosen)>=limit:
                if mandatory:
                    raise DocumentError('필수 자료를 모두 담기에 분석 용량이 부족합니다. 자료를 나누거나 필수 범위를 조정해 주세요.')
                return False
            chosen.append({**source,'document_name':meta['name'],'metadata':meta,'selection_relevance':relevance})
            ids.add(source['id']); used+=len(source['text'])
            return True
        # Reserve capacity for EVERY required document before optional selection.
        for meta in manifest:
            if meta['required']:
                rows = by_doc[meta['id']]
                fitting = next((r for r in rows if len(r[1]['text'])<=budget-used),None)
                if fitting is None:
                    raise DocumentError('필수 자료의 본문을 분석에 포함할 수 없습니다. 자료 범위를 조정해 주세요.')
                add(fitting,mandatory=True)
        for row in sorted(ranked,key=lambda row:(-PRIORITIES[row[2]['priority']],-row[0],row[1]['id'])):
            add(row)
        if not chosen:
            # Even unrelated inputs must be reviewed as inputs, not invented as evidence.
            fallback=next((rows[0] for rows in by_doc.values() if rows),None)
            if fallback is None:raise DocumentError('분석 가능한 본문 근거가 없습니다.')
            add(fallback,mandatory=True)
        return chosen

    def crop(self, key, region):
        doc = self.get(key)
        sid = key+'-'+region
        source = next((s for s in doc['sources'] if s['id']==sid),None)
        if source and doc['format'] in ('xlsx','xls'):
            from excel_evidence import render
            cached=self.root/(key+'-'+region+'-native-v2.png')
            if not cached.exists():
                try: data=render((self.root/(key+'.bin')).read_bytes(),source,doc['format'])
                except Exception as error: raise DocumentError(str(error)) from error
                cached.write_bytes(data)
            return cached.read_bytes()
        if not source or doc['format']!='pdf':
            raise DocumentError('원문 영역을 찾을 수 없습니다.')
        with fitz.open(stream=(self.root/(key+'.bin')).read_bytes(),filetype='pdf') as pdf:
            return pdf[source['page']-1].get_pixmap(matrix=fitz.Matrix(2,2),clip=fitz.Rect(source['bbox']),alpha=False).tobytes('png')
